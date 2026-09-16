import docx

doc = docx.Document()

# Title Page / Header Style
doc.add_heading('APPENDIX A', level=1)
doc.add_heading('PROGRAM LISTINGS', level=2)

# Section 1
doc.add_heading('1. Production Settings Configuration (settings.py)', level=3)
code_settings = """""" + \
"""from pathlib import Path
import os

BASE_DIR = Path(__file__).resolve().parent.parent

SECRET_KEY = os.environ.get('SECRET_KEY', 'django-insecure-fallback-key')
DEBUG = os.environ.get('DEBUG', 'False') == 'True'

INSTALLED_APPS = [
    'django.contrib.admin',
    'django.contrib.auth',
    'django.contrib.contenttypes',
    'django.contrib.sessions',
    'django.contrib.messages',
    'django.contrib.staticfiles',
    'advisory',
]

MIDDLEWARE = [
    'django.middleware.security.SecurityMiddleware',
    'whitenoise.middleware.WhiteNoiseMiddleware',
    'django.contrib.sessions.middleware.SessionMiddleware',
    'django.middleware.common.CommonMiddleware',
    'django.middleware.csrf.CsrfViewMiddleware',
    'django.contrib.auth.middleware.AuthenticationMiddleware',
    'django.contrib.messages.middleware.MessageMiddleware',
    'django.middleware.clickjacking.XFrameOptionsMiddleware',
]

ROOT_URLCONF = 'core.urls'
WSGI_APPLICATION = 'core.wsgi.application'

DATABASES = {
    'default': {
        'ENGINE': 'django.db.backends.sqlite3',
        'NAME': BASE_DIR / 'db.sqlite3',
    }
}

STATIC_URL = 'static/'
STATIC_ROOT = BASE_DIR / 'staticfiles'
STATICFILES_STORAGE = 'whitenoise.storage.CompressedManifestStaticFilesStorage'

ALLOWED_HOSTS = [
    host.strip() 
    for host in os.environ.get('ALLOWED_HOSTS', '127.0.0.1,localhost,testserver,.railway.app').split(',') 
    if host.strip()
]

CSRF_TRUSTED_ORIGINS = [
    origin.strip() 
    for origin in os.environ.get('CSRF_TRUSTED_ORIGINS', 'https://*.railway.app').split(',') 
    if origin.strip()
]"""
p1 = doc.add_paragraph(code_settings)
p1.style = 'Normal'
p1.runs[0].font.name = 'Consolas'
p1.runs[0].font.size = docx.shared.Pt(9.5)

# Section 2
doc.add_heading('2. Core Application Models (advisory/models.py)', level=3)
code_models = """from django.db import models
from django.contrib.auth.models import User

class AdvisoryRecord(models.Model):
    STATUS_CHOICES = [
        ('PENDING', 'Pending Review'),
        ('APPROVED', 'Approved'),
        ('REJECTED', 'Rejected'),
    ]

    student = models.ForeignKey(User, on_delete=models.CASCADE, related_name='advisory_submissions')
    title = models.CharField(max_length=200)
    description = models.TextField()
    status = models.CharField(max_length=20, choices=STATUS_CHOICES, default='PENDING')
    feedback = models.TextField(blank=True, null=True)
    created_at = models.DateTimeField(auto_now_add=True)
    updated_at = models.DateTimeField(auto_now=True)

    def __str__(self):
        return f"{self.title} - {self.student.username} ({self.status})" """
p2 = doc.add_paragraph(code_models)
p2.runs[0].font.name = 'Consolas'
p2.runs[0].font.size = docx.shared.Pt(9.5)

# Section 3
doc.add_heading('3. Application Views (advisory/views.py)', level=3)
code_views = """from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from .models import AdvisoryRecord

@login_required
def dashboard_view(request):
    \"\"\"Renders the main advisory dashboard based on user role.\"\"\"
    user = request.user
    if user.is_staff:
        records = AdvisoryRecord.objects.all().order_by('-created_at')
    else:
        records = AdvisoryRecord.objects.filter(student=user).order_by('-created_at')
    
    context = {
        'records': records,
        'is_staff': user.is_staff,
    }
    return render(request, 'advisory/dashboard.html', context)"""
p3 = doc.add_paragraph(code_views)
p3.runs[0].font.name = 'Consolas'
p3.runs[0].font.size = docx.shared.Pt(9.5)

# Section 4
doc.add_heading('4. WSGI Application Entry Point (core/wsgi.py)', level=3)
code_wsgi = """import os
from django.core.wsgi import get_wsgi_application

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'core.settings')
application = get_wsgi_application()"""
p4 = doc.add_paragraph(code_wsgi)
p4.runs[0].font.name = 'Consolas'
p4.runs[0].font.size = docx.shared.Pt(9.5)

doc.add_page_break()

# Appendix B
doc.add_heading('APPENDIX B', level=1)
doc.add_heading('SAMPLE OUTPUTS', level=2)
doc.add_paragraph('The figures presented in this appendix illustrate key graphical user interfaces and system outputs of the deployed Django hybrid_system application running on Railway.')

doc.add_paragraph("""
-----------------------------------------------------------------
                      HYBRID SYSTEM                          
                   Sign in to Continue                       
                                                             
   Username: [ _________________________________________ ]   
   Password: [ ***************************************** ]   
                                                             
                          [ SIGN IN ]                        
-----------------------------------------------------------------
Figure B.1: System Authentication / Login Interface. Provides secure access control configured with CSRF protection and token sessions.
""")

doc.add_paragraph("""
-----------------------------------------------------------------
Hybrid Admin      Dashboard | Records | Settings    [Admin] 
-----------------------------------------------------------------
Welcome, Administrator                                      
                                                             
  +--------------------+  +-------------------------------+  
  | Total Records: 24  |  | Pending Review: 5             |  |
  +--------------------+  +-------------------------------+  
                                                             
  Recent Submissions:                                        
  - Advisory Request #102 | Student: Emerie | Status: PENDING
  - Advisory Request #101 | Student: John   | Status: APPROVED
-----------------------------------------------------------------
Figure B.2: Main Advisory Dashboard View. Displays administrative summaries, active tracking counts, and status indicators.
""")

doc.add_paragraph("""
-----------------------------------------------------------------
Advisory Submission Form                    [Active Session]
-----------------------------------------------------------------
Title:      [ Academic Progression Review                 ] 
Description:[ Requesting clearance and advising validation  ] 
            [ for the upcoming semester curriculum requirements.
                                                             
                          [ SUBMIT REQUEST ]                 
-----------------------------------------------------------------
Figure B.3: Student Advisory Submission Interface. Enables users to file structured requests securely into the database layer.
""")

doc.save('Appendices_Corrected.docx')
print("Successfully generated Appendices_Corrected.docx!")